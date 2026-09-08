"""Bounded CSV and Excel statement parser."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.core.exceptions import AppError
from app.core.upload_config import (
    PARSER_MAX_COLUMNS,
    PARSER_MAX_LINES,
    PARSER_MAX_ROWS,
    PARSER_MAX_SHEETS,
)
from app.models.enums import TransactionType
from app.services.parsers.base import ParsedLine, ParseResult


class CsvExcelParser:
    name = "csv_excel_parser"

    def parse_bytes(
        self,
        data: bytes,
        filename: str,
        *,
        issuer: str | None = None,
        max_rows: int = PARSER_MAX_ROWS,
        max_columns: int = PARSER_MAX_COLUMNS,
        max_sheets: int = PARSER_MAX_SHEETS,
        max_lines: int = PARSER_MAX_LINES,
    ) -> ParseResult:
        del issuer
        try:
            if filename.lower().endswith((".csv", ".tsv")):
                rows = self._read_delimited(
                    data,
                    delimiter="\t" if filename.lower().endswith(".tsv") else ",",
                    max_rows=max_rows,
                    max_columns=max_columns,
                )
            else:
                rows = self._read_xlsx(
                    data,
                    max_rows=max_rows,
                    max_columns=max_columns,
                    max_sheets=max_sheets,
                )
        except AppError:
            raise
        except Exception as exc:
            raise AppError(
                "The tabular statement is malformed or unreadable", code="tabular_read_failed"
            ) from exc

        if not rows:
            raise AppError("The statement has no tabular rows", code="missing_columns")

        columns = {key.lower().strip(): key for key in rows[0] if key.strip()}
        date_col = next((key for key in columns if "date" in key or "time" in key), None)
        desc_col = next(
            (
                key
                for key in columns
                if "desc" in key or "merchant" in key or "narrative" in key or "particulars" in key
            ),
            None,
        )
        amt_col = next((key for key in columns if "amount" in key or "value" in key), None)
        debit_col = next((key for key in columns if "debit" in key or "withdrawal" in key), None)
        credit_col = next((key for key in columns if "credit" in key or "deposit" in key), None)
        type_col = next((key for key in columns if "type" in key or "cr/dr" in key), None)

        if not date_col or not desc_col:
            raise AppError("Missing required columns: Date, Description", code="missing_columns")
        if not amt_col and not (debit_col or credit_col):
            raise AppError("Missing Amount or Debit/Credit columns", code="missing_columns")

        lines: list[ParsedLine] = []
        for row in rows[1:]:
            raw_date = _value(row, columns[date_col])
            if not raw_date:
                continue
            occurred_at = _parse_datetime(raw_date)
            merchant = _value(row, columns[desc_col]) or "Unknown"
            amount = Decimal("0")
            tx_type = TransactionType.PURCHASE

            if amt_col:
                amount = _amount(_value(row, columns[amt_col]))
                if (
                    type_col
                    and _value(row, columns[type_col]).lower()
                    in {"cr", "credit", "deposit", "payment"}
                ) or _value(row, columns[amt_col]).lstrip().startswith("-"):
                    tx_type = TransactionType.PAYMENT_TO_ISSUER
            elif debit_col and _value(row, columns[debit_col]):
                amount = _amount(_value(row, columns[debit_col]))
            elif credit_col and _value(row, columns[credit_col]):
                amount = _amount(_value(row, columns[credit_col]))
                tx_type = TransactionType.PAYMENT_TO_ISSUER

            if amount <= 0:
                continue
            if len(lines) >= max_lines:
                raise AppError(
                    "The statement exceeds the parser line limit", code="parser_resource_limit"
                )
            lines.append(
                ParsedLine(
                    occurred_at=occurred_at,
                    merchant=merchant[:255],
                    amount_paise=int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
                    proposed_type=tx_type,
                    raw={key: _json_value(value) for key, value in row.items()},
                )
            )
        return ParseResult(lines=lines, parser_name=self.name)

    @staticmethod
    def _read_delimited(
        data: bytes,
        *,
        delimiter: str,
        max_rows: int,
        max_columns: int,
    ) -> list[dict[str, object]]:
        csv.field_size_limit(64 * 1024)
        reader = csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter=delimiter)
        try:
            header = next(reader)
        except StopIteration:
            return []
        if len(header) > max_columns:
            raise AppError(
                "The statement exceeds the parser column limit", code="parser_resource_limit"
            )
        normalized_header = [
            value.strip() or f"column_{index + 1}" for index, value in enumerate(header)
        ]
        rows: list[dict[str, object]] = [
            {key: value for key, value in zip(normalized_header, header, strict=False)}
        ]
        for row_number, values in enumerate(reader, start=1):
            if row_number > max_rows:
                raise AppError(
                    "The statement exceeds the parser row limit", code="parser_resource_limit"
                )
            if len(values) > max_columns:
                raise AppError(
                    "The statement exceeds the parser column limit", code="parser_resource_limit"
                )
            rows.append(
                {
                    key: values[index] if index < len(values) else None
                    for index, key in enumerate(normalized_header)
                }
            )
        return rows

    @staticmethod
    def _read_xlsx(
        data: bytes,
        *,
        max_rows: int,
        max_columns: int,
        max_sheets: int,
    ) -> list[dict[str, object]]:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if len(workbook.sheetnames) > max_sheets:
                raise AppError(
                    "The spreadsheet exceeds the parser sheet limit", code="parser_resource_limit"
                )
            worksheet = workbook[workbook.sheetnames[0]]
            iterator = worksheet.iter_rows(values_only=True)
            try:
                header_values = next(iterator)
            except StopIteration:
                return []
            header = list(header_values)
            if len(header) > max_columns:
                raise AppError(
                    "The statement exceeds the parser column limit", code="parser_resource_limit"
                )
            normalized_header = [
                str(value).strip()
                if value is not None and str(value).strip()
                else f"column_{index + 1}"
                for index, value in enumerate(header)
            ]
            rows: list[dict[str, object]] = [
                {
                    key: value
                    for key, value in zip(normalized_header, header_values, strict=False)
                }
            ]
            for row_number, values in enumerate(iterator, start=1):
                if row_number > max_rows:
                    raise AppError(
                        "The statement exceeds the parser row limit", code="parser_resource_limit"
                    )
                values_list = list(values)
                if len(values_list) > max_columns:
                    raise AppError(
                        "The statement exceeds the parser column limit",
                        code="parser_resource_limit",
                    )
                rows.append(
                    {
                        key: values_list[index] if index < len(values_list) else None
                        for index, key in enumerate(normalized_header)
                    }
                )
            return rows
        finally:
            workbook.close()


def _value(row: dict[str, object], key: str) -> str:
    value = row.get(key)
    if value is None:
        return ""
    return str(value).strip()


def _parse_datetime(value: str) -> datetime | None:
    for fmt in (None, "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%m-%d-%Y"):
        try:
            return datetime.fromisoformat(value) if fmt is None else datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def _amount(value: str) -> Decimal:
    cleaned = value.replace(",", "").replace("₹", "").strip()
    if not cleaned:
        return Decimal("0")
    try:
        return abs(Decimal(cleaned))
    except InvalidOperation:
        return Decimal("0")


def _json_value(value: Any) -> object:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)
